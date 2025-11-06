"""
Generador de reportes en formato Excel
"""
from typing import Any, Dict, List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

from .base_report import BaseReportGenerator


class ExcelReportGenerator(BaseReportGenerator):
    """
    Generador de reportes en formato Excel usando openpyxl
    """
    
    def __init__(self, titulo: str = "Reporte", categoria: str = "general", parametros: Dict[str, Any] = None):
        super().__init__(titulo, categoria, parametros)
        self.wb = None
        self.ws = None
    
    @staticmethod
    def generar_reporte(titulo: str, datos: Any, categoria: str) -> str:
        """
        Método estático de conveniencia para generar un reporte Excel
        
        Args:
            titulo: Título del reporte
            datos: Datos del reporte (lista de diccionarios o dict estructurado)
            categoria: Categoría del reporte
            
        Returns:
            str: Ruta del archivo generado
        """
        generator = ExcelReportGenerator(titulo=titulo, categoria=categoria)
        
        # Convertir datos si es una lista simple
        if isinstance(datos, list):
            data_estructurada = generator._convertir_lista_a_estructura(datos, categoria)
        else:
            data_estructurada = datos
        
        return generator.generar(data_estructurada)
    
    def _convertir_lista_a_estructura(self, datos: List[Dict], categoria: str) -> Dict:
        """Convierte una lista de diccionarios en la estructura esperada para Excel"""
        if not datos:
            return {
                'hojas': [{
                    'nombre': categoria.capitalize(),
                    'encabezado': {},
                    'tablas': []
                }]
            }
        
        # Obtener las columnas de los datos
        columnas = list(datos[0].keys())
        
        # Crear encabezados de tabla legibles
        headers = [self._formatear_columna(col) for col in columnas]
        
        # Crear filas de datos
        filas = [[self._formatear_valor(row.get(col)) for col in columnas] for row in datos]
        
        # Construir estructura
        return {
            'hojas': [{
                'nombre': categoria.capitalize(),
                'encabezado': {},
                'tablas': [{
                    'titulo': f'Datos de {categoria.capitalize()}',
                    'data': [headers] + filas
                }]
            }]
        }
    
    def _formatear_columna(self, nombre: str) -> str:
        """Convierte nombres de columna en títulos legibles"""
        return ' '.join(word.capitalize() for word in nombre.replace('_', ' ').split())
    
    def _formatear_valor(self, valor: Any) -> str:
        """Formatea un valor para mostrar en Excel"""
        if valor is None:
            return ''
        if isinstance(valor, (datetime, )):
            return valor.strftime("%d/%m/%Y %H:%M")
        return str(valor)
    
    def generar(self, data: Any) -> str:
        """
        Genera un reporte Excel
        
        Args:
            data: Diccionario con la estructura del reporte
                {
                    'hojas': [
                        {
                            'nombre': 'Hoja1',
                            'encabezado': {...},
                            'tablas': [...]
                        }
                    ]
                }
        
        Returns:
            str: Ruta relativa del archivo generado
        """
        filename = self._get_filename('xlsx')
        filepath = self._get_filepath(filename)
        
        # Crear workbook
        self.wb = Workbook()
        
        # Eliminar la hoja por defecto
        if 'Sheet' in self.wb.sheetnames:
            del self.wb['Sheet']
        
        # Procesar hojas
        hojas = data.get('hojas', [{'nombre': 'Reporte', 'tablas': data.get('tablas', [])}])
        
        for hoja_data in hojas:
            self._crear_hoja(hoja_data)
        
        # Guardar archivo
        self.wb.save(filepath)
        
        return self._get_relative_path(filepath)
    
    def _crear_hoja(self, hoja_data: Dict):
        """Crea una hoja en el workbook"""
        nombre_hoja = hoja_data.get('nombre', 'Reporte')
        self.ws = self.wb.create_sheet(nombre_hoja)
        
        fila_actual = 1
        
        # Agregar encabezado
        if hoja_data.get('encabezado'):
            fila_actual = self._agregar_encabezado(hoja_data['encabezado'], fila_actual)
        
        # Agregar tablas
        for tabla_data in hoja_data.get('tablas', []):
            fila_actual = self._agregar_tabla(tabla_data, fila_actual)
            fila_actual += 2  # Espacio entre tablas
    
    def _agregar_encabezado(self, encabezado: Dict, fila_inicio: int) -> int:
        """Agrega el encabezado del reporte"""
        # Título principal
        self.ws.merge_cells(f'A{fila_inicio}:F{fila_inicio}')
        celda_titulo = self.ws[f'A{fila_inicio}']
        celda_titulo.value = self.titulo
        celda_titulo.font = Font(size=16, bold=True, color='1E3A8A')
        celda_titulo.alignment = Alignment(horizontal='center', vertical='center')
        fila_inicio += 1
        
        # Fecha de generación
        self.ws.merge_cells(f'A{fila_inicio}:F{fila_inicio}')
        celda_fecha = self.ws[f'A{fila_inicio}']
        fecha_actual = datetime.now().strftime("%d/%m/%Y %H:%M")
        celda_fecha.value = f"Fecha de generación: {fecha_actual}"
        celda_fecha.font = Font(size=10, italic=True)
        celda_fecha.alignment = Alignment(horizontal='right')
        fila_inicio += 1
        
        # Período (si existe)
        if encabezado.get('fecha_inicio') or encabezado.get('fecha_fin'):
            self.ws.merge_cells(f'A{fila_inicio}:F{fila_inicio}')
            celda_periodo = self.ws[f'A{fila_inicio}']
            fecha_inicio = self._format_fecha(encabezado.get('fecha_inicio', ''))
            fecha_fin = self._format_fecha(encabezado.get('fecha_fin', ''))
            celda_periodo.value = f"Período: {fecha_inicio} - {fecha_fin}"
            celda_periodo.font = Font(size=10, italic=True)
            celda_periodo.alignment = Alignment(horizontal='right')
            fila_inicio += 1
        
        return fila_inicio + 2  # Espacio después del encabezado
    
    def _agregar_tabla(self, tabla_data: Dict, fila_inicio: int) -> int:
        """Agrega una tabla al reporte"""
        fila_actual = fila_inicio
        
        # Título de la tabla
        if tabla_data.get('titulo'):
            self.ws.merge_cells(f'A{fila_actual}:F{fila_actual}')
            celda_titulo = self.ws[f'A{fila_actual}']
            celda_titulo.value = tabla_data['titulo']
            celda_titulo.font = Font(size=12, bold=True, color='3B82F6')
            celda_titulo.alignment = Alignment(horizontal='left')
            fila_actual += 1
        
        # Datos de la tabla
        data = tabla_data.get('data', [])
        if not data:
            return fila_actual
        
        # Estilos
        estilo_encabezado = {
            'font': Font(bold=True, color='FFFFFF', size=11),
            'fill': PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid'),
            'alignment': Alignment(horizontal='center', vertical='center'),
            'border': Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
        }
        
        estilo_celda = {
            'alignment': Alignment(horizontal='left', vertical='center'),
            'border': Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
        }
        
        # Agregar encabezados
        for col_idx, valor in enumerate(data[0], start=1):
            celda = self.ws.cell(row=fila_actual, column=col_idx)
            celda.value = valor
            celda.font = estilo_encabezado['font']
            celda.fill = estilo_encabezado['fill']
            celda.alignment = estilo_encabezado['alignment']
            celda.border = estilo_encabezado['border']
        
        fila_actual += 1
        
        # Agregar datos
        for fila_data in data[1:]:
            for col_idx, valor in enumerate(fila_data, start=1):
                celda = self.ws.cell(row=fila_actual, column=col_idx)
                celda.value = valor
                celda.alignment = estilo_celda['alignment']
                celda.border = estilo_celda['border']
                
                # Alternar colores de fila
                if fila_actual % 2 == 0:
                    celda.fill = PatternFill(start_color='F3F4F6', end_color='F3F4F6', fill_type='solid')
            
            fila_actual += 1
        
        # Ajustar ancho de columnas
        for col_idx in range(1, len(data[0]) + 1):
            column_letter = get_column_letter(col_idx)
            max_length = 0
            for row in self.ws[column_letter]:
                if row.value:
                    max_length = max(max_length, len(str(row.value)))
            adjusted_width = min(max_length + 2, 50)
            self.ws.column_dimensions[column_letter].width = adjusted_width
        
        return fila_actual
