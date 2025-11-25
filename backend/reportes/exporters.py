"""
GeneradorArchivos: Servicio para generar archivos PDF y Excel.
Utiliza reportlab para PDF y openpyxl para Excel.
"""
from io import BytesIO
from typing import Dict, List, Any
from datetime import datetime
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


class PDFExporter:
    """Exportador a PDF usando ReportLab"""
    
    def generar(self, datos_reporte: Dict) -> BytesIO:
        """
        Generar PDF del reporte
        
        Args:
            datos_reporte: Diccionario con datos, columnas, título, etc.
        
        Returns:
            BytesIO con el contenido del PDF
        """
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        elementos = []
        styles = getSampleStyleSheet()
        
        # Estilo para título
        style_titulo = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#1e40af'),
            spaceAfter=12,
            alignment=1  # Centrado
        )
        
        # Estilo para subtítulo
        style_subtitulo = ParagraphStyle(
            'CustomSubtitle',
            parent=styles['Normal'],
            fontSize=10,
            textColor=colors.grey,
            spaceAfter=20,
            alignment=1
        )
        
        # Agregar título
        titulo = Paragraph(datos_reporte['titulo'], style_titulo)
        elementos.append(titulo)
        
        # Agregar subtítulo
        if datos_reporte.get('subtitulo'):
            subtitulo = Paragraph(datos_reporte['subtitulo'], style_subtitulo)
            elementos.append(subtitulo)
        
        # Agregar información de generación
        info = Paragraph(
            f"Generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M')} | "
            f"Total de registros: {datos_reporte['total_registros']}",
            styles['Normal']
        )
        elementos.append(info)
        elementos.append(Spacer(1, 0.3 * inch))
        
        # Crear tabla de datos
        if datos_reporte['datos']:
            # Preparar datos para la tabla
            columnas = datos_reporte['columnas']
            datos_tabla = [columnas]  # Primera fila: encabezados
            
            # Agregar filas de datos
            for registro in datos_reporte['datos']:
                fila = []
                for col in columnas:
                    # Mapear nombres de columnas a claves del diccionario
                    key = self._normalizar_key(col)
                    valor = registro.get(key, '')
                    
                    # Formatear valores
                    if isinstance(valor, float):
                        if 'total' in key or 'monto' in key or 'precio' in key or 'ingreso' in key:
                            valor = f"Bs {valor:,.2f}"
                        else:
                            valor = f"{valor:,.2f}"
                    
                    fila.append(str(valor))
                datos_tabla.append(fila)
            
            # Crear tabla
            tabla = Table(datos_tabla)
            
            # Estilo de la tabla
            tabla.setStyle(TableStyle([
                # Encabezado
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                
                # Datos
                ('BACKGROUND', (0, 1), (-1, -1), colors.white),
                ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
                ('ALIGN', (0, 1), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 9),
                ('TOPPADDING', (0, 1), (-1, -1), 6),
                ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
                
                # Bordes
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('BOX', (0, 0), (-1, -1), 2, colors.HexColor('#1e40af')),
                
                # Alternar colores de fila
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f3f4f6')]),
            ]))
            
            elementos.append(tabla)
        else:
            elementos.append(Paragraph("No hay datos para mostrar", styles['Normal']))
        
        # Construir PDF
        doc.build(elementos)
        buffer.seek(0)
        return buffer
    
    def _normalizar_key(self, columna: str) -> str:
        """Convertir nombre de columna a key del diccionario"""
        normalized = columna.lower()
        normalized = normalized.replace(' de ', '_').replace(' del ', '_').replace(' la ', '_')
        normalized = normalized.replace(' ', '_')
        normalized = normalized.replace('á', 'a').replace('é', 'e').replace('í', 'i').replace('ó', 'o').replace('ú', 'u')
        return normalized
    
    def generar_multiple(self, reportes: list) -> BytesIO:
        """
        Generar PDF con múltiples reportes en el mismo documento
        
        Args:
            reportes: Lista de diccionarios con datos de reportes
        
        Returns:
            BytesIO con el contenido del PDF
        """
        buffer = BytesIO()
        
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )
        
        elementos = []
        styles = getSampleStyleSheet()
        
        # Agregar cada reporte
        for idx, datos_reporte in enumerate(reportes):
            if idx > 0:
                elementos.append(PageBreak())
            
            titulo_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=18,
                textColor=colors.HexColor('#1e40af'),
                spaceAfter=12,
                alignment=1
            )
            
            titulo = Paragraph(datos_reporte['titulo'], titulo_style)
            elementos.append(titulo)
            
            if datos_reporte.get('subtitulo'):
                subtitulo_style = ParagraphStyle(
                    'CustomSubtitle',
                    parent=styles['Normal'],
                    fontSize=10,
                    textColor=colors.grey,
                    spaceAfter=8,
                    alignment=1
                )
                subtitulo = Paragraph(datos_reporte['subtitulo'], subtitulo_style)
                elementos.append(subtitulo)
            
            info_style = ParagraphStyle(
                'CustomInfo',
                parent=styles['Normal'],
                fontSize=9,
                spaceAfter=20,
                alignment=1
            )
            info = Paragraph(
                f"Generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M')} | Total de registros: {datos_reporte['total_registros']}",
                info_style
            )
            elementos.append(info)
            
            if len(reportes) > 1:
                indicador_style = ParagraphStyle(
                    'Indicador',
                    parent=styles['Normal'],
                    fontSize=10,
                    textColor=colors.HexColor('#1e40af'),
                    spaceAfter=10,
                    alignment=1,
                    fontName='Helvetica-Bold'
                )
                indicador = Paragraph(f"Reporte {idx + 1} de {len(reportes)}", indicador_style)
                elementos.append(indicador)
            
            if datos_reporte['datos']:
                datos_tabla = [datos_reporte['columnas']]
                
                for registro in datos_reporte['datos']:
                    fila = []
                    for columna in datos_reporte['columnas']:
                        key = self._normalizar_key(columna)
                        valor = registro.get(key, '')
                        
                        if isinstance(valor, (int, float)):
                            if 'total' in key or 'monto' in key or 'precio' in key or 'ingreso' in key:
                                valor = f"Bs {valor:,.2f}"
                            else:
                                valor = f"{valor:,.2f}"
                        
                        fila.append(str(valor))
                    
                    datos_tabla.append(fila)
                
                tabla = Table(datos_tabla)
                
                tabla.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 10),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.white),
                    ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
                    ('ALIGN', (0, 1), (-1, -1), 'LEFT'),
                    ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                    ('FONTSIZE', (0, 1), (-1, -1), 9),
                    ('TOPPADDING', (0, 1), (-1, -1), 6),
                    ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
                    ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                    ('BOX', (0, 0), (-1, -1), 2, colors.HexColor('#1e40af')),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f3f4f6')]),
                ]))
                
                elementos.append(tabla)
            else:
                elementos.append(Paragraph("No hay datos para mostrar", styles['Normal']))
        
        doc.build(elementos)
        buffer.seek(0)
        return buffer


class ExcelExporter:
    """Exportador a Excel usando openpyxl"""
    
    def generar(self, datos_reporte: Dict) -> BytesIO:
        """
        Generar Excel del reporte
        
        Args:
            datos_reporte: Diccionario con datos, columnas, título, etc.
        
        Returns:
            BytesIO con el contenido del Excel
        """
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Reporte"
        
        titulo_font = Font(size=16, bold=True, color="1e40af")
        subtitulo_font = Font(size=10, color="666666")
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        fila_actual = 1
        
        col_count = len(datos_reporte['columnas'])
        sheet.merge_cells(f'A{fila_actual}:{get_column_letter(col_count)}{fila_actual}')
        cell_titulo = sheet[f'A{fila_actual}']
        cell_titulo.value = datos_reporte['titulo']
        cell_titulo.font = titulo_font
        cell_titulo.alignment = Alignment(horizontal='center')
        fila_actual += 1
        
        if datos_reporte.get('subtitulo'):
            sheet.merge_cells(f'A{fila_actual}:{get_column_letter(col_count)}{fila_actual}')
            cell_subtitulo = sheet[f'A{fila_actual}']
            cell_subtitulo.value = datos_reporte['subtitulo']
            cell_subtitulo.font = subtitulo_font
            cell_subtitulo.alignment = Alignment(horizontal='center')
            fila_actual += 1
        
        sheet.merge_cells(f'A{fila_actual}:{get_column_letter(col_count)}{fila_actual}')
        cell_info = sheet[f'A{fila_actual}']
        cell_info.value = f"Generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M')} | Total de registros: {datos_reporte['total_registros']}"
        cell_info.font = Font(size=9)
        cell_info.alignment = Alignment(horizontal='center')
        fila_actual += 2
        
        for col_idx, columna in enumerate(datos_reporte['columnas'], start=1):
            cell = sheet.cell(row=fila_actual, column=col_idx)
            cell.value = columna
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = border
        
        fila_actual += 1
        
        for registro in datos_reporte['datos']:
            for col_idx, columna in enumerate(datos_reporte['columnas'], start=1):
                key = self._normalizar_key(columna)
                valor = registro.get(key, '')
                
                cell = sheet.cell(row=fila_actual, column=col_idx)
                
                if isinstance(valor, str) and valor.startswith('Bs '):
                    try:
                        cell.value = float(valor.replace('Bs ', '').replace(',', ''))
                        cell.number_format = '"Bs "#,##0.00'
                    except:
                        cell.value = valor
                elif isinstance(valor, (int, float)):
                    cell.value = valor
                    if 'total' in key or 'monto' in key or 'precio' in key or 'ingreso' in key:
                        cell.number_format = '"Bs "#,##0.00'
                    else:
                        cell.number_format = '#,##0.00'
                else:
                    cell.value = str(valor)
                
                cell.border = border
                
                if fila_actual % 2 == 0:
                    cell.fill = PatternFill(start_color="f3f4f6", end_color="f3f4f6", fill_type="solid")
            
            fila_actual += 1
        
        for col_idx in range(1, col_count + 1):
            col_letter = get_column_letter(col_idx)
            max_length = 0
            
            for cell in sheet[col_letter]:
                if isinstance(cell, MergedCell):
                    continue
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            sheet.column_dimensions[col_letter].width = adjusted_width
        
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer
    
    def generar_multiple(self, reportes: list) -> BytesIO:
        """
        Generar Excel con múltiples reportes en hojas separadas
        
        Args:
            reportes: Lista de diccionarios con datos de reportes
        
        Returns:
            BytesIO con el contenido del Excel
        """
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        
        titulo_font = Font(size=16, bold=True, color="1e40af")
        subtitulo_font = Font(size=10, color="666666")
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for idx, datos_reporte in enumerate(reportes):
            titulo_hoja = datos_reporte.get('titulo', f'Reporte {idx + 1}')[:31]
            sheet = workbook.create_sheet(title=titulo_hoja)
            
            col_count = len(datos_reporte['columnas'])
            fila_actual = 1
            
            if len(reportes) > 1:
                sheet.merge_cells(f'A{fila_actual}:{get_column_letter(col_count)}{fila_actual}')
                cell_indicador = sheet[f'A{fila_actual}']
                cell_indicador.value = f"Reporte {idx + 1} de {len(reportes)}"
                cell_indicador.font = Font(size=12, bold=True, color="1e40af")
                cell_indicador.alignment = Alignment(horizontal='center')
                fila_actual += 1
            
            sheet.merge_cells(f'A{fila_actual}:{get_column_letter(col_count)}{fila_actual}')
            cell_titulo = sheet[f'A{fila_actual}']
            cell_titulo.value = datos_reporte['titulo']
            cell_titulo.font = titulo_font
            cell_titulo.alignment = Alignment(horizontal='center')
            fila_actual += 1
            
            if datos_reporte.get('subtitulo'):
                sheet.merge_cells(f'A{fila_actual}:{get_column_letter(col_count)}{fila_actual}')
                cell_subtitulo = sheet[f'A{fila_actual}']
                cell_subtitulo.value = datos_reporte['subtitulo']
                cell_subtitulo.font = subtitulo_font
                cell_subtitulo.alignment = Alignment(horizontal='center')
                fila_actual += 1
            
            sheet.merge_cells(f'A{fila_actual}:{get_column_letter(col_count)}{fila_actual}')
            cell_info = sheet[f'A{fila_actual}']
            cell_info.value = f"Generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M')} | Total de registros: {datos_reporte['total_registros']}"
            cell_info.font = Font(size=9)
            cell_info.alignment = Alignment(horizontal='center')
            fila_actual += 2
            
            for col_idx, columna in enumerate(datos_reporte['columnas'], start=1):
                cell = sheet.cell(row=fila_actual, column=col_idx)
                cell.value = columna
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center')
                cell.border = border
            
            fila_actual += 1
            
            for registro in datos_reporte['datos']:
                for col_idx, columna in enumerate(datos_reporte['columnas'], start=1):
                    key = self._normalizar_key(columna)
                    valor = registro.get(key, '')
                    
                    cell = sheet.cell(row=fila_actual, column=col_idx)
                    
                    if isinstance(valor, str) and valor.startswith('Bs '):
                        try:
                            cell.value = float(valor.replace('Bs ', '').replace(',', ''))
                            cell.number_format = '"Bs "#,##0.00'
                        except:
                            cell.value = valor
                    elif isinstance(valor, (int, float)):
                        cell.value = valor
                        if 'total' in key or 'monto' in key or 'precio' in key or 'ingreso' in key:
                            cell.number_format = '"Bs "#,##0.00'
                        else:
                            cell.number_format = '#,##0.00'
                    else:
                        cell.value = str(valor)
                    
                    cell.border = border
                    
                    if fila_actual % 2 == 0:
                        cell.fill = PatternFill(start_color="f3f4f6", end_color="f3f4f6", fill_type="solid")
                
                fila_actual += 1
            
            for col_idx in range(1, col_count + 1):
                col_letter = get_column_letter(col_idx)
                max_length = 0
                
                for cell in sheet[col_letter]:
                    if isinstance(cell, MergedCell):
                        continue
                    try:
                        if cell.value and len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                
                adjusted_width = min(max_length + 2, 50)
                sheet.column_dimensions[col_letter].width = adjusted_width
        
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer
    
    def _normalizar_key(self, columna: str) -> str:
        """Convertir nombre de columna a key del diccionario"""
        normalized = columna.lower()
        normalized = normalized.replace(' de ', '_').replace(' del ', '_').replace(' la ', '_')
        normalized = normalized.replace(' ', '_')
        normalized = normalized.replace('á', 'a').replace('é', 'e').replace('í', 'i').replace('ó', 'o').replace('ú', 'u')
        return normalized

